"""
DOCX extractor preserving document structure, headings, lists, and tables.
"""
import os
from typing import List, Dict, Any
from docx import Document
from extractors.models import DocumentContent, ContentBlock

class DocxExtractor:
    @staticmethod
    def extract(file_path: str) -> DocumentContent:
        doc = Document(file_path)
        file_name = os.path.basename(file_path)
        blocks: List[ContentBlock] = []
        raw_text_parts: List[str] = []
        block_counter = 0

        # 1. Process paragraphs
        for p_idx, p in enumerate(doc.paragraphs):
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name.lower() if p.style else ""
            block_type = "paragraph"
            heading_level = 2

            if "heading 1" in style_name or "title" in style_name:
                block_type = "heading"
                heading_level = 1
            elif "heading 2" in style_name:
                block_type = "heading"
                heading_level = 2
            elif "heading 3" in style_name:
                block_type = "heading"
                heading_level = 3
            elif "list" in style_name or "bullet" in style_name or text.startswith(("•", "-", "*")):
                block_type = "bullet"

            block = ContentBlock(
                block_id=f"docx_p_{block_counter}",
                block_type=block_type,
                text=p.text,
                location={"paragraph_index": p_idx, "section": "body"},
                style_metadata={"style_name": p.style.name if p.style else "", "level": heading_level}
            )
            blocks.append(block)
            raw_text_parts.append(p.text)
            block_counter += 1

        # 2. Process tables
        for t_idx, table in enumerate(doc.tables):
            for r_idx, row in enumerate(table.rows):
                row_cells_text = [cell.text.strip() for cell in row.cells]
                row_str = " | ".join(row_cells_text)
                
                block = ContentBlock(
                    block_id=f"docx_tbl_{t_idx}_r_{r_idx}",
                    block_type="table_row",
                    text=row_str,
                    location={"table_index": t_idx, "row": r_idx, "cell_count": len(row.cells)},
                    style_metadata={"is_header": (r_idx == 0)}
                )
                blocks.append(block)
                raw_text_parts.append(row_str)
                block_counter += 1

        full_raw_text = "\n".join(raw_text_parts)
        metadata = {
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
            "source_type": "docx"
        }

        return DocumentContent(
            file_name=file_name,
            file_type="docx",
            raw_text=full_raw_text,
            blocks=blocks,
            metadata=metadata
        )
