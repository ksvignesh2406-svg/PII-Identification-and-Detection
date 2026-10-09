"""
Generate realistic Cadence financial sample artifacts for Case Study 2:
- .pptx (Cadence Wealth Management Review)
- .pdf (Cadence Credit Risk Policy & Loan Evidence)
- .txt (Cadence Internal Audit Incident Memo)
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

samples_dir = os.path.dirname(os.path.abspath(__file__))

def create_sample_pptx():
    pptx_path = os.path.join(samples_dir, "cadence_wealth_portfolio_review.pptx")
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(5.625) # 16:9

    # Slide 1: Title
    blank_slide_layout = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(blank_slide_layout)
    tb = s1.shapes.add_textbox(Inches(1), Inches(1.5), Inches(8), Inches(2))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Cadence Financial Services"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(14, 43, 92)
    p2 = tf.add_paragraph()
    p2.text = "Q3 High Net Worth Client Portfolio & KYC Due Diligence Review"
    p2.font.size = Pt(20)
    p2.font.color.rgb = RGBColor(220, 80, 40)
    p3 = tf.add_paragraph()
    p3.text = "Author: Marcus Sterling (EMP-88392) | Division: Private Wealth Management"
    p3.font.size = Pt(13)
    p3.font.italic = True

    # Slide 2: Table of High Net Worth Clients
    s2 = prs.slides.add_slide(blank_slide_layout)
    tb = s2.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.4), Inches(0.8))
    p = tb.text_frame.paragraphs[0]
    p.text = "Client Portfolios & Identity Verification Records"
    p.font.size = Pt(22)
    p.font.bold = True

    rows, cols = 4, 5
    table_shape = s2.shapes.add_table(rows, cols, Inches(0.8), Inches(1.3), Inches(8.4), Inches(2.2))
    table = table_shape.table
    headers = ["Client Name", "Tax ID / SSN", "Account Number", "Contact Email", "Phone"]
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        if cell.text_frame.paragraphs and cell.text_frame.paragraphs[0].runs:
            cell.text_frame.paragraphs[0].runs[0].font.bold = True
            cell.text_frame.paragraphs[0].runs[0].font.size = Pt(12)

    data = [
        ["Eleanor Vance", "142-88-9031", "ACCT-902148291", "eleanor.vance@vancetech.com", "+1 (212) 555-0194"],
        ["David K. Richardson", "873-12-4490", "ACCT-449182304", "drichardson@richardsoncap.com", "+1 (646) 555-8821"],
        ["Sarah Jenkins-Wu", "309-54-1182", "ACCT-773829105", "sarah.j.wu@apexglobal.net", "+1 (917) 555-3419"]
    ]
    for r_idx, row_data in enumerate(data, start=1):
        for c_idx, val in enumerate(row_data):
            cell = table.cell(r_idx, c_idx)
            cell.text = val
            if cell.text_frame.paragraphs and cell.text_frame.paragraphs[0].runs:
                cell.text_frame.paragraphs[0].runs[0].font.size = Pt(11)

    # Slide 3: Advisory Note & Wire Instructions
    s3 = prs.slides.add_slide(blank_slide_layout)
    tb = s3.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(8.4), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "CONFIDENTIAL: Wire Instructions & Settlement Officer Notes"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = RGBColor(14, 43, 92)

    notes = [
        ("Beneficiary Name: ", "Jonathan Hayes (Cadence Escrow Services)"),
        ("Officer Contact: ", "jonathan.hayes@cadence-corp.com | Direct: +1 (212) 555-7733"),
        ("Beneficiary Address: ", "380 Madison Avenue, 18th Floor, New York, NY 10017"),
        ("Bank Routing Number (ABA): ", "021000021 (JPMorgan Chase NY)"),
        ("Cadence Settlement Account: ", "883920194012"),
        ("SWIFT / BIC: ", "CHASUS33XXX"),
        ("Note: ", "Customer SSN 982-11-4092 confirmed against NY State Driver License # NY-7749201. Do not share raw identifiers across unencrypted channels.")
    ]
    for title, detail in notes:
        p = tf.add_paragraph()
        run1 = p.add_run()
        run1.text = title
        run1.font.bold = True
        run1.font.size = Pt(13)
        run2 = p.add_run()
        run2.text = detail
        run2.font.size = Pt(13)

    prs.save(pptx_path)
    print("Created PPTX:", pptx_path)

def create_sample_pdf():
    pdf_path = os.path.join(samples_dir, "cadence_credit_policy_evaluation.pdf")
    doc = SimpleDocTemplate(pdf_path, pagesize=letter, rightMargin=40, leftMargin=40, topMargin=40, bottomMargin=40)
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#0e2b5c')
    )
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#dc5028')
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#1f2937')
    )

    story = []
    story.append(Paragraph("CADENCE FINANCIAL SERVICES — CREDIT RISK EVALUATION", title_style))
    story.append(Spacer(1, 10))
    story.append(Paragraph("Underwriting Assessment Artifact # CR-2026-904 | Classification: Highly Confidential", h2_style))
    story.append(Spacer(1, 10))

    intro_text = (
        "This evaluation artifact records the commercial credit assessment conducted by Cadence Underwriting Team. "
        "The following documentation includes borrower identity proofs, banking histories, and credit bureau scores "
        "required for corporate loan underwriting compliance under regulatory guidelines."
    )
    story.append(Paragraph(intro_text, body_style))
    story.append(Spacer(1, 12))

    story.append(Paragraph("1. Primary Applicant & Guarantor Verification", h2_style))
    applicant_details = (
        "<b>Primary Borrower:</b> Robert Anthony Sterling<br/>"
        "<b>Date of Birth:</b> 12/04/1981<br/>"
        "<b>Social Security Number (SSN):</b> 458-92-3019<br/>"
        "<b>Residential Address:</b> 742 Evergreen Terrace, Apt 4B, White Plains, NY 10601<br/>"
        "<b>Personal Email:</b> robert.sterling@sterling-holdings.com<br/>"
        "<b>Cell Phone:</b> +1 914-555-0182<br/>"
        "<b>US Passport Number:</b> C09823145 (Expires: 08/2031)<br/>"
        "<b>State Driver License:</b> NY-DL902184910"
    )
    story.append(Paragraph(applicant_details, body_style))
    story.append(Spacer(1, 12))

    story.append(Paragraph("2. Financial Accounts & Banking Relationship", h2_style))
    table_data = [
        ["Institution", "Account Type", "Account Number", "Routing Number", "Balance USD"],
        ["Cadence Commercial", "Operating Checking", "670192830194", "026009593", "$1,450,200"],
        ["Citibank N.A.", "Escrow Reserve", "440192817290", "021000089", "$850,000"],
        ["Corporate Card", "Visa Infinite", "4532 9012 3456 7810", "N/A", "$42,100"]
    ]
    t = Table(table_data, colWidths=[110, 110, 120, 100, 90])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#0e2b5c')),
        ('TEXTCOLOR', (0, 0), (-1, 0), colors.whitesmoke),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, 0), 9),
        ('BOTTOMPADDING', (0, 0), (-1, 0), 6),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#cbd5e1')),
        ('BACKGROUND', (0, 1), (-1, -1), colors.HexColor('#f8fafc')),
        ('FONTSIZE', (0, 1), (-1, -1), 9),
    ]))
    story.append(t)
    story.append(Spacer(1, 14))

    story.append(Paragraph("3. Sign-off & Audit Trail", h2_style))
    sign_off = (
        "Reviewed and verified by Senior Underwriter: <b>Emily Watson</b> (Employee ID: <b>EMP-55109</b>).<br/>"
        "Direct email: <b>emily.watson@cadence-corp.com</b> | Department: Risk Governance Division.<br/>"
        "Underwriter confirms that all customer identifiers have been verified. Under Cadence AI Policy 2026-R4, "
        "no raw customer PII may be processed by downstream generative AI models without prior tokenization."
    )
    story.append(Paragraph(sign_off, body_style))

    doc.build(story)
    print("Created PDF:", pdf_path)

def create_sample_txt():
    txt_path = os.path.join(samples_dir, "cadence_internal_audit_memo.txt")
    content = """CADENCE FINANCIAL SERVICES - INTERNAL AUDIT REPORT
MEMORANDUM | STRICTLY CONFIDENTIAL
DATE: September 28, 2026
TO: Cadence TPRM & AI Governance Committee
FROM: Daniel Miller (daniel.miller@cadence-corp.com, Emp ID: EMP-11029)
SUBJECT: Security Incident Review & PII Containment

During the audit of external vendor sub-processor interfaces, the following contractor and employee credentials were observed in transmission logs:

1. Lead Contractor: Alexander Gomez
   Email: a.gomez@fintech-partners.io
   Phone: +1 650-555-9204
   Passport ID: P88291044
   Assigned Bank Wire Account: 992014827103
   Aadhaar (Overseas Specialist): 9012 3456 7890

2. Senior Risk Manager: Dr. Priya Raman
   Email: priya.raman@cadence-corp.com
   Phone: +91 98840 12345
   PAN Card: ABCPR1234K
   Employee ID: EMP-44021

3. Incident Context:
   Vendor API submitted customer records without sanitization. Sensitive data included SSN 219-40-5812 and credit card 5425 2334 8901 2345.
   Recommendation: Immediate implementation of Optiv Automated PII Gateway before LLM ingestion.
"""
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(content)
    print("Created TXT:", txt_path)

if __name__ == "__main__":
    create_sample_pptx()
    create_sample_pdf()
    create_sample_txt()
