from docx import Document

from .analyzer import analyze as _analyze


def redact_paragraph(paragraph):
    full_text = "".join(run.text for run in paragraph.runs)

    if not full_text:
        return

    results = _analyze(full_text)

    if not results:
        return

    results = sorted(results, key=lambda x: x.start, reverse=True)

    for result in results:
        replacement = f"<{result.entity_type}>"
        current_position = 0

        for run in paragraph.runs:
            run_start = current_position
            run_end = current_position + len(run.text)
            current_position = run_end

            if result.end <= run_start or result.start >= run_end:
                continue

            if result.start >= run_start and result.end <= run_end:
                start_in_run = result.start - run_start
                end_in_run = result.end - run_start
                run.text = (
                    run.text[:start_in_run]
                    + replacement
                    + run.text[end_in_run:]
                )
                break

            if result.start >= run_start and result.start < run_end:
                start_in_run = result.start - run_start
                run.text = run.text[:start_in_run] + replacement
            elif result.start < run_start < result.end:
                if result.end >= run_end:
                    run.text = ""
                else:
                    end_in_run = result.end - run_start
                    run.text = run.text[end_in_run:]


def _redact_table(table):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                redact_paragraph(paragraph)
            for nested_table in cell.tables:
                _redact_table(nested_table)


def process_docx(input_file, output_file):
    doc = Document(input_file)

    for paragraph in doc.paragraphs:
        redact_paragraph(paragraph)

    for table in doc.tables:
        _redact_table(table)

    processed_parts = set()
    for section in doc.sections:
        for part_name in (
            "header",
            "first_page_header",
            "even_page_header",
            "footer",
            "first_page_footer",
            "even_page_footer",
        ):
            part = getattr(section, part_name)
            if part.part.partname in processed_parts:
                continue
            processed_parts.add(part.part.partname)
            for paragraph in part.paragraphs:
                redact_paragraph(paragraph)
            for table in part.tables:
                _redact_table(table)

    doc.save(output_file)


if __name__ == "__main__":
    input_file = "uploads/pii_redaction_test.docx"
    output_file = "outputs/redacted1.docx"
    process_docx(input_file, output_file)
    print("Redaction complete!")
    print("Output:", output_file)